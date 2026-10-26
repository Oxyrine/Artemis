import hashlib
import json
import shutil
from pathlib import Path

import pytest
from openpyxl import load_workbook

from artemis.db import Db, Invoice, RemittanceEmail
from artemis.ingest.load_all import LoadError, load_all
from data.generator import generate

ROOT = Path(__file__).resolve().parents[2]


def digest(root: Path) -> dict[str, str]:
    return {str(p.relative_to(root)): hashlib.sha256(p.read_bytes()).hexdigest()
            for p in sorted(root.rglob("*")) if p.is_file()}


def build_into(tmp: Path, monkeypatch) -> dict:
    monkeypatch.setattr(generate, "ENV", tmp / "env")
    monkeypatch.setattr(generate, "TRUTH", tmp / "cases_truth.json")
    return generate.write(generate.build())


def test_generator_is_byte_identical_and_matches_what_is_committed(tmp_path, monkeypatch):
    a, b = tmp_path / "a", tmp_path / "b"
    build_into(a, monkeypatch)
    build_into(b, monkeypatch)
    assert digest(a) == digest(b)  # re-running gives byte-identical files
    assert digest(a / "env") == digest(ROOT / "data" / "env")  # and the committed environment is that output
    assert (a / "cases_truth.json").read_bytes() == (ROOT / "data" / "cases_truth.json").read_bytes()


def test_answer_key_split_and_hero(tmp_path, monkeypatch):
    truth = build_into(tmp_path, monkeypatch)
    split = {}
    for t in truth.values():
        split[t["status"]] = split.get(t["status"], 0) + 1
    assert split == {"AUTO_CLOSED": 10, "VERIFIED": 16, "PROBABLE": 6, "CONFLICTING": 4, "INSUFFICIENT": 4}
    assert sum(len(t["receipts"]) for t in truth.values()) == 42 and len(truth) == 40
    hero = truth["CASE-0007"]
    assert hero["receipts"] == ["UTR839201"] and hero["adjustments"] == [{"kind": "tds", "amount_paise": 2_600_000}]
    assert [a["invoice_id"] for a in hero["allocations"]] == ["INV-1042", "INV-1043"]
    assert truth["CASE-0008"]["status"] == "PROBABLE" and truth["CASE-0009"]["status"] == "CONFLICTING"
    assert sum(t["is_red_team"] for t in truth.values()) == 5  # R1 + R2 + R3 (two cases) + R4


@pytest.mark.parametrize("case,customer", [("case_UTR839201", "CUST-014"), ("case_UTR839202", "CUST-015"),
                                           ("case_UTR839203", "CUST-016")])
def test_lane_b_fixtures_agree_with_the_real_environment(case, customer, real):
    """Lane B's hero fixtures must never drift from the generated environment they stand in for."""
    fx = json.loads((ROOT / "fixtures" / f"{case}.json").read_text(encoding="utf-8"))
    with real.core.db.tx() as s:
        for inv in fx["invoices"]:
            row = s.get(Invoice, inv["id"])
            assert (row.customer_id, row.taxable_value, row.gst_amount, row.gross_amount) == (
                customer, inv["taxable_value"], inv["gst_amount"], inv["gross_amount"])


def test_load_all_loads_everything_as_integer_paise(real):
    with real.core.db.tx() as s:
        assert s.query(Invoice).count() == 70
        inv = s.get(Invoice, "INV-1042")
        assert isinstance(inv.gross_amount, int) and inv.gross_amount == 118_000_000
        eml = s.query(RemittanceEmail).first()
        assert eml.body_text and "<" not in eml.body_text and eml.raw_path.startswith("emails/")  # plain text, no HTML
        assert eml.customer_id is not None


def env_copy(tmp: Path) -> Path:
    dst = tmp / "env"
    shutil.copytree(ROOT / "data" / "env", dst)
    return dst


def test_every_validation_breach_fails_loudly_and_loads_nothing(tmp_path):
    env = env_copy(tmp_path)
    bank = env / "bank_statement.csv"
    lines = bank.read_text(encoding="utf-8").splitlines()
    bank.write_text("\n".join(lines + [lines[1]]) + "\n", encoding="utf-8")  # repeat the first UTR
    db = Db("sqlite://")
    db.reset()
    with pytest.raises(LoadError, match="duplicate UTR"):
        load_all(db, env)
    with db.tx() as s:
        assert s.query(Invoice).count() == 0  # one transaction: nothing half-loaded

    env2 = env_copy(tmp_path / "x")
    wb = load_workbook(env2 / "erp_invoices.xlsx")
    wb.active.cell(row=2, column=7).value = "1.00"  # gross no longer equals taxable + GST
    wb.save(env2 / "erp_invoices.xlsx")
    db.reset()
    with pytest.raises(LoadError, match="gross"):
        load_all(db, env2)

    env3 = env_copy(tmp_path / "y")
    tax = env3 / "tax_credit_extract.csv"
    rows = tax.read_text(encoding="utf-8").splitlines()
    head, first = rows[0].split(","), rows[1].split(",")
    first[head.index("quarter")] = ""
    tax.write_text("\n".join([rows[0], ",".join(first)] + rows[2:]) + "\n", encoding="utf-8")
    db.reset()
    with pytest.raises(LoadError, match="quarter"):
        load_all(db, env3)
